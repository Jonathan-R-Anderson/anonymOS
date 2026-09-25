from pptx import Presentation


def analyze_ppt(file_path):
    print(f"Analyzing {file_path}...")
    try:
        prs = Presentation(file_path)
    except Exception as e:
        print(f"Error loading PPT: {e}")
        return

    print(f"Number of slides: {len(prs.slides)}")

    for i, slide in enumerate(prs.slides):
        print(f"\n--- Slide {i + 1} ---")
        for shape in slide.shapes:
            if shape.is_placeholder:
                print(
                    f"Placeholder: shape.name='{shape.name}', shape.placeholder_format.idx={shape.placeholder_format.idx}",
                )
                if shape.has_text_frame:
                    print(f"  Text: {shape.text.strip()}")
            elif shape.has_text_frame:
                print(f"Text Box: shape.name='{shape.name}'")
                print(f"  Text: {shape.text.strip()}")


if __name__ == "__main__":
    analyze_ppt("/mnt/work/projects/spectre/docs/Project PPT Format.pptx")
