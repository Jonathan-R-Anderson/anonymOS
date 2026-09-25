from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


def add_bullets(slide, items, top_inch=2.0, font_size=18):
    left = Inches(1.0)
    top = Inches(top_inch)
    width = Inches(8.0)
    height = Inches(4.5)

    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True

    for i, item in enumerate(items):
        p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
        p.text = item
        p.level = 0
        p.font.size = Pt(font_size)
        p.font.name = "Calibri"
        # Add some space between paragraphs
        p.space_after = Pt(14)


def main():
    template_path = "/mnt/work/projects/spectre/docs/Project PPT Format.pptx"
    output_path = "/mnt/work/projects/spectre/docs/Project_Spectre_Submission.pptx"

    prs = Presentation(template_path)

    # Slide 1 (Placeholders)
    s1 = prs.slides[0]
    title_ph = s1.placeholders[0]
    title_ph.text = "Project Spectre: Behavioral Host Intrusion Detection System\nGTU SBTP - 2026 (Subject Code: DI05000011)"

    info_ph = s1.placeholders[1]
    info_ph.text = (
        "Student Name: Aayush Bankar\n"
        "Enrollment no.: 246230316006\n"
        "College Name: Government Polytechnic Gandhinagar\n"
        "Branch and Year: Information Technology, 2026\n"
        "Branch code: 16\n"
        "Gender: Male\n"
        "Registered Mail ID: aayushbankar42@gmail.com\n"
        "Contact No.: +91 6351400725"
    )
    # Fix font size for slide 1
    for p in info_ph.text_frame.paragraphs:
        p.font.size = Pt(16)

    # Slide 2: What did you learn?
    add_bullets(
        prs.slides[1],
        [
            "Systems Architecture & Security: Designed and implemented a low-level, behavioral Host Intrusion Detection System (HIDS).",
            "Advanced Threat Modeling: Shifted the security paradigm from static file analysis to dynamic execution pattern modeling.",
            "Agile Engineering (SDLC): Executed a 10+ stage incremental Software Development Life Cycle to continuously ship functional versions.",
            "Modern Tech Stack Integration: Gained hands-on experience integrating system telemetry (psutil), graph algorithms (NetworkX), malware engines (YARA), and full-stack APIs (FastAPI, Next.js).",
        ],
    )

    # Slide 3: Purpose
    add_bullets(
        prs.slides[2],
        [
            "The Problem: Traditional antivirus relies heavily on static signatures ('Is this file known?'), leaving systems vulnerable to zero-day threats and 'living off the land' techniques.",
            "The Solution (Spectre): Built to analyze the grammar of a computer system, asking: 'Does this sequence of actions make sense on this machine?'",
            "The Vision: To engineer a lightweight, local-first detection system that prioritizes execution relationships to spot anomalies, while remaining fully explainable and resource-efficient.",
        ],
    )

    # Slide 4: Objectives
    add_bullets(
        prs.slides[3],
        [
            "Behavioral Monitoring: Track process ancestry, file I/O operations, and network socket connections in real-time.",
            "Intelligent Detection Engine: Implement weighted threat scoring, JSON-configurable rules, and dynamic anomaly thresholds.",
            "Active Containment: Engineer mitigation modules to instantly quarantine, freeze (SIGSTOP), or terminate (SIGKILL) malicious process trees.",
            "Contextual Enrichment: Integrate MITRE ATT&CK framework mapping and YARA signature scanning to provide deep threat context.",
        ],
    )

    # Slide 5: Structure
    add_bullets(
        prs.slides[4],
        [
            "1. Telemetry Sensing: Extract real-time OS-level data on processes and resources.",
            "2. Graph Construction: Normalize events into a rolling-memory, sliding-window process-resource graph.",
            "3. Behavioral Evaluation: Run the detection engine to evaluate active execution chains against configured threat rules.",
            "4. Alerting & Mitigation: Generate human-readable alerts, map to MITRE ATT&CK, and trigger active containment.",
            "5. Visualization: Expose telemetry via a FastAPI REST API and visualize live threats on a Next.js web dashboard.",
        ],
    )

    # Slide 6: GitHub
    # Find the github link text box and replace it
    for shape in prs.slides[5].shapes:
        if shape.has_text_frame and "github.com" in shape.text:
            shape.text = "https://github.com/Aayushbankar/spectre.git"
            for p in shape.text_frame.paragraphs:
                p.font.size = Pt(24)
                p.alignment = PP_ALIGN.CENTER
            break

    # Slide 7: Advantages and Limitations
    add_bullets(
        prs.slides[6],
        [
            "ADVANTAGES:",
            "• Zero-Day Resilience: Detects novel attacks by analyzing behavior rather than relying on known signatures.",
            "• Local-First & Explainable: Operates entirely offline with strict data privacy.",
            "• Ultra-Lightweight: Engineered for minimal footprint (Targeting CPU <5% and RAM <200MB).",
            "",
            "LIMITATIONS:",
            "• Polling Constraints: Application-level polling can occasionally miss micro-lived processes.",
            "• Processing Overhead: High-activity environments require careful threshold tuning to prevent bottlenecks.",
        ],
        font_size=16,
    )

    # Slide 8: Challenges
    add_bullets(
        prs.slides[7],
        [
            "State Management: Safely handling short-lived processes and PID recycling during ancestry tree construction to prevent tracking errors.",
            "Memory Optimization: Implementing an efficient sliding-window graph algorithm to automatically prune stale events and prevent memory leaks.",
            "False Positive Tuning: Accurately mapping complex behaviors to MITRE ATT&CK techniques without flagging legitimate administrative workflows.",
            "Safe Mitigation: Designing a containment module capable of neutralizing threat trees without accidentally disrupting critical OS services.",
        ],
    )

    # Slide 9: Conclusion
    add_bullets(
        prs.slides[8],
        [
            "Project Spectre proves that a behavior-first, relationship-driven approach to endpoint security is both highly viable and highly effective.",
            "Moving away from static characteristics allows for deep, contextual visibility into execution sequences.",
            "The system establishes a robust, explainable, and machine-learning-ready foundation for next-generation Endpoint Detection and Response (EDR) platforms.",
        ],
    )

    # Save the file
    prs.save(output_path)
    print(f"Successfully generated {output_path}")


if __name__ == "__main__":
    main()
